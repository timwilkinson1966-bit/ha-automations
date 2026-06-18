#!/usr/bin/env python3
"""
Builds a starter compliance workbook (compliance_workbook_template.xlsx).

This is a *source data* template for the Transport Manager to fill in. The
compliance_check.py tool reads workbooks shaped like this and produces the
RAG report + briefing. Edit this script (or the workbook) to match how you
actually want to track things.

Run:  python make_template.py
"""

from __future__ import annotations

import datetime as dt

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TODAY = dt.date.today()

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13)
NOTE_FONT = Font(italic=True, size=9, color="808080")
DATE_FMT = "DD/MM/YYYY"
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def d(offset_days: int) -> dt.date:
    """A date relative to today — used only for the example rows."""
    return TODAY + dt.timedelta(days=offset_days)


def add_sheet(wb, title, intro, headers, rows, date_cols, widths):
    ws = wb.create_sheet(title)
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = intro
    ws["A2"].font = NOTE_FONT

    header_row = 4
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for r_off, row in enumerate(rows, start=header_row + 1):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r_off, column=c, value=val)
            cell.border = BORDER
            if c in date_cols and isinstance(val, dt.date):
                cell.number_format = DATE_FMT

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{header_row + 1}"
    ws.row_dimensions[header_row].height = 30
    return ws


def build() -> str:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # drop default sheet

    # --- Vehicles -----------------------------------------------------------
    add_sheet(
        wb,
        "Vehicles",
        "One row per vehicle. Dates drive the RAG report — fill the due dates you track.",
        ["Reg", "Fleet No", "Make / Model", "MOT Due", "Road Tax Due",
         "PMI Due", "Tacho Calibration Due", "LOLER Due", "Insurance Due", "Notes"],
        [
            ["AB12 CDE", "V01", "DAF CF",   d(-6),  d(40),  d(4),   d(120), d(300), d(200), "PMI due this week"],
            ["FG34 HIJ", "V02", "Volvo FH", d(95),  d(10),  d(22),  d(60),  d(15),  d(180), ""],
            ["KL56 MNO", "V03", "Scania R", d(210), d(150), d(180), d(240), d(330), d(95),  "All current"],
        ],
        date_cols={4, 5, 6, 7, 8, 9},
        widths=[12, 9, 16, 13, 13, 12, 16, 12, 13, 28],
    )

    # --- Trailers -----------------------------------------------------------
    add_sheet(
        wb,
        "Trailers",
        "One row per trailer. LOLER applies if it has a tail-lift / lifting equipment.",
        ["Trailer No", "Type", "MOT Due", "PMI Due", "LOLER Due", "Notes"],
        [
            ["TR-101", "Curtainsider", d(30),  d(8),   d(-3),  "LOLER overdue"],
            ["TR-102", "Fridge",       d(120), d(45),  d(60),  ""],
        ],
        date_cols={3, 4, 5},
        widths=[12, 16, 13, 12, 12, 28],
    )

    # --- Drivers ------------------------------------------------------------
    add_sheet(
        wb,
        "Drivers",
        "One row per driver. Licence check = your periodic DVLA check date.",
        ["Driver Name", "Employee No", "Licence Check Due", "Licence Expiry",
         "CPC Expiry", "Digi Tacho Card Expiry", "Medical Due", "Notes"],
        [
            ["John Smith", "D01", d(-2),  d(800), d(420), d(50),  d(900), "Licence check overdue"],
            ["Jane Doe",   "D02", d(12),  d(600), d(12),  d(200), d(365), "CPC due soon"],
            ["Raj Patel",  "D03", d(150), d(950), d(700), d(400), d(120), ""],
        ],
        date_cols={3, 4, 5, 6, 7},
        widths=[16, 12, 17, 14, 13, 20, 13, 26],
    )

    # --- Operator Licence ---------------------------------------------------
    add_sheet(
        wb,
        "Operator Licence",
        "Standing O-licence items. One row per item; 'Due Date' is what gets checked.",
        ["Item", "Reference", "Due Date", "Notes"],
        [
            ["O-Licence continuation",      "OB1234567", d(75),  "5-yearly continuation"],
            ["Financial standing review",   "—",         d(25),  "Evidence funds available"],
            ["Maintenance contract review", "—",         d(-10), "Review with provider"],
        ],
        date_cols={3},
        widths=[30, 14, 13, 30],
    )

    out = "compliance_workbook_template.xlsx"
    wb.save(out)
    return out


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
